"""Document processor — extracts structured text from multiple document formats."""

import html
from html.parser import HTMLParser
import logging
from dataclasses import dataclass
from pathlib import Path
import tempfile

logger = logging.getLogger(__name__)


@dataclass
class ProcessedDocument:
    """Result of document processing — structured text with metadata."""

    source_file: str
    file_name: str
    document_type: str
    content_markdown: str
    page_count: int = 1
    metadata: dict | None = None


class _HTMLToMarkdownParser(HTMLParser):
    """Lightweight parser to convert HTML email body into readable Markdown/text."""

    def __init__(self):
        super().__init__()
        self._pieces: list[str] = []
        self._skip_depth: int = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in ("script", "style", "head", "meta", "link"):
            self._skip_depth += 1
            return
        if self._skip_depth > 0:
            return

        if tag in ("p", "div", "tr", "section", "article"):
            self._pieces.append("\n\n")
        elif tag == "br":
            self._pieces.append("\n")
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(tag[1])
            self._pieces.append(f"\n\n{'#' * level} ")
        elif tag == "li":
            self._pieces.append("\n- ")
        elif tag == "hr":
            self._pieces.append("\n\n---\n\n")
        elif tag in ("td", "th"):
            self._pieces.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in ("script", "style", "head", "meta", "link"):
            if self._skip_depth > 0:
                self._skip_depth -= 1
            return
        if self._skip_depth > 0:
            return

        if tag in ("p", "div", "section", "article", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
            self._pieces.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            self._pieces.append(data)

    def get_text(self) -> str:
        text = "".join(self._pieces)
        text = html.unescape(text)
        lines = [line.strip() for line in text.splitlines()]
        result: list[str] = []
        blank = False
        for line in lines:
            if line:
                result.append(line)
                blank = False
            elif not blank:
                result.append("")
                blank = True
        return "\n".join(result).strip()


class DocumentProcessor:
    """Extract structured text from PDF, Word, Excel, Image, and Text files.

    Uses free, open-source libraries for all processing. Designed so this
    class can be swapped with an Azure Document Intelligence implementation.
    """

    def process(self, file_path: str | Path) -> ProcessedDocument:
        """Process a document file and return structured Markdown content."""
        path = Path(file_path)
        ext = path.suffix.lower()

        if not path.exists():
            raise FileNotFoundError(f"Document not found: {path}")

        logger.info(f"Processing {path.name} ({ext})")

        match ext:
            case ".pdf":
                return self._process_pdf(path)
            case ".docx":
                return self._process_docx(path)
            case ".xlsx":
                return self._process_xlsx(path)
            case ".png" | ".jpg" | ".jpeg":
                return self._process_image(path)
            case ".txt" | ".md":
                return self._process_text(path)
            case ".eml":
                return self._process_eml(path)
            case _:
                raise ValueError(f"Unsupported file format: {ext}")

    def _process_pdf(self, path: Path) -> ProcessedDocument:
        """Extract content from PDF using Docling (primary) or pypdf (fallback)."""
        content = ""
        page_count = 0

        # Try Docling first for layout-aware extraction
        try:
            from docling.document_converter import DocumentConverter

            converter = DocumentConverter()
            result = converter.convert(str(path))
            content = result.document.export_to_markdown()
            # Estimate page count from the docling result
            page_count = max(1, content.count("\n---\n") + 1)
            logger.info(f"PDF processed with Docling: {path.name}")
        except Exception as e:
            logger.warning(f"Docling failed for {path.name}: {e}. Falling back to pypdf.")
            content, page_count = self._process_pdf_fallback(path)

        return ProcessedDocument(
            source_file=str(path),
            file_name=path.name,
            document_type="pdf",
            content_markdown=content.strip(),
            page_count=page_count,
        )

    def _process_pdf_fallback(self, path: Path) -> tuple[str, int]:
        """Fallback PDF processing using pypdf."""
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"## Page {i}\n\n{text.strip()}")

        return "\n\n".join(pages), len(reader.pages)

    def _process_docx(self, path: Path) -> ProcessedDocument:
        """Extract content from Word documents using python-docx."""
        from docx import Document as DocxDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = DocxDocument(str(path))
        sections: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = (para.style.name or "").lower()

            # Convert heading styles to Markdown headings
            if "heading 1" in style_name:
                sections.append(f"# {text}")
            elif "heading 2" in style_name:
                sections.append(f"## {text}")
            elif "heading 3" in style_name:
                sections.append(f"### {text}")
            elif "heading" in style_name:
                sections.append(f"#### {text}")
            elif "list" in style_name:
                sections.append(f"- {text}")
            else:
                sections.append(text)

        # Extract tables
        for table in doc.tables:
            table_md = self._table_to_markdown(
                [[cell.text.strip() for cell in row.cells] for row in table.rows]
            )
            sections.append(table_md)

        content = "\n\n".join(sections)
        logger.info(f"DOCX processed: {path.name} ({len(sections)} sections)")

        return ProcessedDocument(
            source_file=str(path),
            file_name=path.name,
            document_type="docx",
            content_markdown=content.strip(),
        )

    def _process_xlsx(self, path: Path) -> ProcessedDocument:
        """Extract content from Excel files using openpyxl (sheet by sheet)."""
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        sheets: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else "" for cell in row]
                if any(cell.strip() for cell in row_data):
                    rows.append(row_data)

            if rows:
                table_md = self._table_to_markdown(rows)
                sheets.append(f"## Sheet: {sheet_name}\n\n{table_md}")

        wb.close()
        content = "\n\n".join(sheets)
        logger.info(f"XLSX processed: {path.name} ({len(sheets)} sheets)")

        return ProcessedDocument(
            source_file=str(path),
            file_name=path.name,
            document_type="xlsx",
            content_markdown=content.strip(),
        )

    def _process_image(self, path: Path) -> ProcessedDocument:
        """Extract text from images using Tesseract OCR."""
        try:
            from PIL import Image
            import pytesseract

            image = Image.open(str(path))
            text = pytesseract.image_to_string(image)
        except ImportError:
            logger.error("pytesseract or Pillow not installed. Cannot process images.")
            text = f"[Image: {path.name} — OCR not available]"
        except Exception as e:
            logger.error(f"OCR failed for {path.name}: {e}")
            text = f"[Image: {path.name} — OCR failed: {e}]"

        content = f"## Image: {path.name}\n\n{text.strip()}"
        logger.info(f"Image processed: {path.name}")

        return ProcessedDocument(
            source_file=str(path),
            file_name=path.name,
            document_type="image",
            content_markdown=content.strip(),
        )

    def _process_text(self, path: Path) -> ProcessedDocument:
        """Read plain text or Markdown files directly."""
        content = path.read_text(encoding="utf-8", errors="replace")
        logger.info(f"Text file read: {path.name}")

        return ProcessedDocument(
            source_file=str(path),
            file_name=path.name,
            document_type="text",
            content_markdown=content.strip(),
        )

    @staticmethod
    def _table_to_markdown(rows: list[list[str]]) -> str:
        """Convert a 2D list of strings to a Markdown table."""
        if not rows:
            return ""

        # Use first row as header
        header = rows[0]
        separator = ["---"] * len(header)
        lines = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(separator) + " |",
        ]
        for row in rows[1:]:
            # Pad or trim row to match header width
            padded = row + [""] * (len(header) - len(row))
            lines.append("| " + " | ".join(padded[: len(header)]) + " |")

        return "\n".join(lines)

    def _process_eml(self, path: Path) -> ProcessedDocument:
        """Extract content from .eml email messages including headers, body, and attachments."""
        import email
        from email import policy

        try:
            with open(path, "rb") as f:
                msg = email.message_from_binary_file(f, policy=policy.default)
        except Exception as e:
            logger.error(f"Failed to read EML file {path.name}: {e}")
            raise

        subject = str(msg.get("subject", "")).strip() or "(No Subject)"
        from_addr = str(msg.get("from", "")).strip()
        to_addr = str(msg.get("to", "")).strip()
        cc_addr = str(msg.get("cc", "")).strip()
        date_str = str(msg.get("date", "")).strip()

        # 1. Extract email body
        body_text = self._extract_email_body(msg)

        # 2. Extract attachments
        attachments = self._extract_email_attachments(msg)

        # 3. Assemble sections
        sections: list[str] = []

        # Document Header & Metadata
        sections.append(f"# Email: {subject}")
        details = [
            "## Email Details",
            f"- **Subject**: {subject}",
            f"- **From**: {from_addr or 'Unknown'}",
            f"- **To**: {to_addr or 'Unknown'}",
        ]
        if cc_addr:
            details.append(f"- **Cc**: {cc_addr}")
        if date_str:
            details.append(f"- **Date**: {date_str}")
        details.append(f"- **Attachments Count**: {len(attachments)}")
        sections.append("\n".join(details))

        # Email Body
        if body_text:
            sections.append(f"## Email Body\n\n{body_text}")
        else:
            sections.append("## Email Body\n\n*(No text content in email body)*")

        # Attachments summary & contents
        total_pages = 1
        if attachments:
            att_summary_lines = ["## Attachments Summary"]
            for idx, att in enumerate(attachments, 1):
                size_str = self._format_size(att["size"])
                att_summary_lines.append(
                    f"- **Attachment {idx}**: `{att['filename']}` ({att['content_type']}, {size_str})"
                )
            sections.append("\n".join(att_summary_lines))

            # Process each attachment
            for idx, att in enumerate(attachments, 1):
                att_name = att["filename"]
                att_bytes = att["bytes"]
                att_type = att["content_type"]
                size_str = self._format_size(att["size"])

                att_content, att_pages = self._process_attachment_content(
                    att_name, att_bytes, att_type
                )
                total_pages += att_pages

                att_section = [
                    f"## Attachment: {att_name}",
                    f"**File**: `{att_name}` | **Type**: {att_type} | **Size**: {size_str}\n",
                    att_content,
                ]
                sections.append("\n\n".join(att_section))

        content_markdown = "\n\n".join(sections).strip()
        logger.info(
            f"EML processed: {path.name} (Subject: '{subject[:40]}...', "
            f"{len(attachments)} attachments, {len(content_markdown)} chars)"
        )

        return ProcessedDocument(
            source_file=str(path),
            file_name=path.name,
            document_type="eml",
            content_markdown=content_markdown,
            page_count=max(1, total_pages),
            metadata={
                "subject": subject,
                "from": from_addr,
                "to": to_addr,
                "date": date_str,
                "attachment_count": len(attachments),
                "attachments": [a["filename"] for a in attachments],
            },
        )

    def _extract_email_body(self, msg) -> str:
        """Extract body text from email message, preferring plain text over HTML."""
        body_part = msg.get_body(preferencelist=("plain", "html"))
        if body_part:
            content_type = body_part.get_content_type()
            try:
                content = body_part.get_content()
                if content_type == "text/html":
                    return self._html_to_markdown(str(content))
                return str(content).strip()
            except Exception as e:
                logger.warning(f"Error reading body from get_body(): {e}")
                payload = body_part.get_payload(decode=True)
                if payload:
                    decoded = payload.decode("utf-8", errors="replace").strip()
                    if content_type == "text/html":
                        return self._html_to_markdown(decoded)
                    return decoded

        # Fallback: walk all parts to locate plain or html body
        html_fallback = ""
        for part in msg.walk():
            if part.get_content_maintype() == "multipart":
                continue
            disp = str(part.get_content_disposition() or "").lower()
            if disp == "attachment":
                continue
            c_type = part.get_content_type()
            if c_type == "text/plain":
                try:
                    return str(part.get_content()).strip()
                except Exception:
                    payload = part.get_payload(decode=True)
                    if payload:
                        return payload.decode("utf-8", errors="replace").strip()
            elif c_type == "text/html" and not html_fallback:
                try:
                    html_fallback = str(part.get_content()).strip()
                except Exception:
                    payload = part.get_payload(decode=True)
                    if payload:
                        html_fallback = payload.decode("utf-8", errors="replace").strip()

        if html_fallback:
            return self._html_to_markdown(html_fallback)

        return ""

    @staticmethod
    def _html_to_markdown(html_content: str) -> str:
        """Convert HTML content into clean readable Markdown/text."""
        if not html_content or not html_content.strip():
            return ""
        parser = _HTMLToMarkdownParser()
        try:
            parser.feed(html_content)
            return parser.get_text()
        except Exception as e:
            logger.warning(f"HTML parsing failed: {e}. Falling back to plain text strip.")
            import re
            cleaned = re.sub(r"<[^>]+>", " ", html_content)
            return html.unescape(cleaned).strip()

    def _extract_email_attachments(self, msg) -> list[dict]:
        """Extract all attachments from an email message."""
        attachments: list[dict] = []
        seen_part_ids = set()

        # 1. Standard attachments via iter_attachments()
        for part in msg.iter_attachments():
            seen_part_ids.add(id(part))
            fn = part.get_filename() or "attachment"
            fn = Path(fn).name
            content_type = part.get_content_type()

            try:
                payload = part.get_content()
                if isinstance(payload, bytes):
                    raw_bytes = payload
                elif isinstance(payload, str):
                    raw_bytes = payload.encode("utf-8")
                else:
                    raw_bytes = part.get_payload(decode=True) or b""
            except Exception:
                raw_bytes = part.get_payload(decode=True) or b""

            attachments.append({
                "filename": fn,
                "content_type": content_type,
                "bytes": raw_bytes,
                "size": len(raw_bytes),
            })

        # 2. Walk parts to find any attachment not captured by iter_attachments
        for part in msg.walk():
            if id(part) in seen_part_ids or part.get_content_maintype() == "multipart":
                continue
            disp = str(part.get_content_disposition() or "").lower()
            fn = part.get_filename()
            if disp == "attachment" or (fn and disp != "inline"):
                fn = Path(fn).name if fn else "attachment"
                content_type = part.get_content_type()
                try:
                    raw_bytes = part.get_payload(decode=True) or b""
                except Exception:
                    raw_bytes = b""

                if raw_bytes:
                    attachments.append({
                        "filename": fn,
                        "content_type": content_type,
                        "bytes": raw_bytes,
                        "size": len(raw_bytes),
                    })
                    seen_part_ids.add(id(part))

        return attachments

    def _process_attachment_content(
        self, filename: str, raw_bytes: bytes, content_type: str
    ) -> tuple[str, int]:
        """Process an email attachment's raw bytes into structured Markdown.

        Returns:
            Tuple of (formatted_markdown, page_count).
        """
        if not raw_bytes:
            return "*(Empty attachment file)*", 1

        ext = Path(filename).suffix.lower()
        supported_doc_exts = {".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".txt", ".md", ".eml"}
        text_exts = {".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log", ".yaml", ".yml"}

        if ext in supported_doc_exts:
            try:
                with tempfile.TemporaryDirectory() as tmp_dir:
                    tmp_file = Path(tmp_dir) / filename
                    tmp_file.write_bytes(raw_bytes)
                    sub_doc = self.process(tmp_file)
                    formatted_content = self._format_attachment_markdown(
                        filename, sub_doc.content_markdown
                    )
                    return formatted_content, max(1, sub_doc.page_count)
            except Exception as e:
                logger.warning(f"Failed to process supported attachment {filename}: {e}")
                return f"*[Error extracting content from {filename}: {e}]*", 1

        elif ext in text_exts or content_type.startswith("text/"):
            try:
                decoded = raw_bytes.decode("utf-8", errors="replace").strip()
                if ext in (".html", ".htm"):
                    decoded = self._html_to_markdown(decoded)
                formatted_content = self._format_attachment_markdown(filename, decoded)
                return formatted_content, 1
            except Exception as e:
                logger.warning(f"Failed to decode text attachment {filename}: {e}")
                return f"*[Error decoding text content from {filename}: {e}]*", 1

        else:
            size_str = self._format_size(len(raw_bytes))
            return (
                f"*[Binary attachment ({content_type}, {size_str}) — content extraction not supported]*",
                1,
            )

    def _format_attachment_markdown(self, filename: str, text: str) -> str:
        """Prefix section headings in extracted attachment text with the attachment name."""
        lines = []
        for line in text.split("\n"):
            stripped = line.strip()
            if stripped.startswith("## Page "):
                page_part = stripped.replace("## Page ", "").strip()
                lines.append(f"### Attachment: {filename} - Page {page_part}")
            elif stripped.startswith("# "):
                heading_text = stripped[2:].strip()
                lines.append(f"### Attachment: {filename} - {heading_text}")
            elif stripped.startswith("## "):
                heading_text = stripped[3:].strip()
                lines.append(f"### Attachment: {filename} - {heading_text}")
            else:
                lines.append(line)
        return "\n".join(lines).strip()

    @staticmethod
    def _format_size(bytes_len: int) -> str:
        """Format bytes into human-readable string."""
        if bytes_len < 1024:
            return f"{bytes_len} B"
        if bytes_len < 1024 * 1024:
            return f"{bytes_len / 1024:.1f} KB"
        return f"{bytes_len / (1024 * 1024):.1f} MB"
